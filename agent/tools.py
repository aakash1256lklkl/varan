"""
Varan tools — JSON-schema function definitions exposed to the AI provider,
plus a dispatcher that executes them on real Office files.
"""
from __future__ import annotations

import ast
import json
import os
import re
from pathlib import Path

from office.registry import Registry
from office.word_live import MatchNotFoundError, WordNotAvailable
from office.excel_live import ExcelNotAvailable
from office.ppt_live import PptNotAvailable, FillerContentError


TOOLS = [
    {
        "type": "function",
        "function": {
            "name": "create_document",
            "description": "Create a new Word (.docx) document.",
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "Output path, e.g. outputs/my_doc.docx"},
                    "title": {"type": "string", "description": "Main title (becomes the title heading)"},
                    "body": {"type": "array", "description": "List of blocks. Each block is a dict with a 'type' of 'heading' (level 1-6), 'paragraph', 'bullet', 'numbered', or 'table'. Paragraphs offer 'text', 'bold', 'italic'. Headings use 'text' and 'level'. Tables use 'headers' (list) and 'rows' (list of lists)."},
                },
                "required": ["path"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "create_workbook",
            "description": "Create a new Excel (.xlsx) workbook.",
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "Output path, e.g. outputs/budget.xlsx"},
                    "sheet": {"type": "string", "description": "First sheet name (default 'Sheet1')"},
                    "data": {"type": "array", "description": "Rows of data. Each row is a dict {'cells': [values...]}. First row can be headers."},
                },
                "required": ["path"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "create_presentation",
            "description": "Create a new PowerPoint (.pptx) presentation.",
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "Output path, e.g. outputs/deck.pptx"},
                    "slides": {"type": "array", "description": "List of slide dicts. Each slide: {'layout': 'title'|'bullets'|'blank', 'title': str, 'subtitle': str, 'bullets': [str], 'table': {'headers': [str], 'rows': [[..]]}, 'chart': {'type': 'bar'|'line'|'pie', 'title': str, 'categories': [str], 'data': [num], 'series_name': str}}"},
                },
                "required": ["path"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "edit_document",
            "description": "Edit an existing Word (.docx) file precisely while preserving the original formatting, tracked changes, comments and styles. If the file is open in Word it edits the LIVE open document (no duplicates); otherwise it writes an '_edited' copy. Replaces/inserts/deletes text without destroying the surrounding bold, italic, fonts or styles. Optionally records real tracked changes (redlines) as a named author. Use read_file first if unsure of exact text.",
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "Path to existing .docx file"},
                    "edits": {"type": "array", "description": "List of edits. Each edit: {'action': 'replace'|'insert_after'|'insert_before'|'delete'|'delete_range', 'match': text to find (or a paragraph anchor like 'P4#d4e5'), 'text': replacement/inserted content, 'count': for 'replace', how many occurrences to replace; set count=-1 to replace EVERY occurrence in the whole document (the reliable way to do a replace-all). 'replace' swaps the matched text in place, keeping its formatting; 'insert_after'/'insert_before' insert new content relative to the matched paragraph; 'delete' removes the matched paragraph; 'delete_range' removes everything from 'match' through 'end_match' (the reliable way to delete a whole page/section — 'end_match' is required, or omit it to delete to the document end). For delete_range you may instead pass 'end_level' (1-6) to end the section just BEFORE the next Heading of that level. The end-anchor matches the LAST occurrence of the text. When the target file is open in Word, edits happen LIVE in the open document."},
                    "table_edits": {"type": "array", "description": "Complex TABLE editing: set the text of specific cells inside existing tables. Each item: {'table': <1-based table index>, 'cell': {'r': row, 'c': col} (0-based) or {'ref': 'A1'}, 'text': new cell text}. Formatting of the cell is preserved. Use this when the user wants a cell in a Word table changed."},
                    "append": {"type": "string", "description": "Text to append at the end of the document"},
                    "content": {"type": "array", "description": "STRUCTURED content to APPEND at the end, rendered as REAL Word paragraphs (not plain text). Each block: {'type': 'heading'|'title'|'paragraph'|'bullet'|'numbered'|'divider', 'level': 1-6 (for heading), 'text': str, 'bold': bool, 'italic': bool, 'font': 'Courier New', 'courier': true}. Use this (NOT create_document) when the user wants a styled section, headings, bullets, italics or a font added to an existing document — it produces real Word styles instead of literal '##'/'-' markdown garbage. 'divider'/'scene_break' inserts a centered '* * *' scene break."},
                    "page": {"type": "integer", "description": "If set, DELETE the entire numeric page (1-based) of the open Word document. Requires the file to be open in Word's interactive window (pagination active). Use when the user says 'delete page N'."},
                    "track_changes": {"type": "boolean", "description": "If true, record the edits as tracked changes (redlines) instead of applying them directly. Default false."},
                    "confirm": {"type": "boolean", "description": "MUST be true to perform a destructive delete ('delete', 'delete_range', or the 'page' parameter), and only AFTER the user has confirmed the exact anchor/section/page to remove. Without confirm=true Varan will refuse the delete and ask which section to remove. Default false."},
                    "remove_blank_pages": {"type": "boolean", "description": "If true, remove every empty paragraph in the document body (the typical cause of 'blank pages' — empty Heading 1 paragraphs used as page separators). This is safe: only paragraphs with NO text are removed; real content is left untouched. Requires no confirm. Use when the user says 'delete blank pages', 'remove empty pages', 'clean up the blank sections'."},
                    "author": {"type": "string", "description": "Author name to attribute tracked changes to (used when track_changes is true). Default 'Varan'."},
                    "inplace": {"type": "boolean", "description": "If true, write the changes back over the original file (no '_edited' copy). Use when editing the user's selected target file so no duplicate copy is created. Default false."}
                },
                "required": ["path"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "edit_workbook",
            "description": "Edit an existing Excel (.xlsx) file and write an '_edited' copy.",
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "Path to existing .xlsx file"},
                    "sheet": {"type": "string", "description": "Target sheet name (default active)"},
                    "writes": {"type": "array", "description": "List of {'cell': 'A1', 'value': ...}"},
                    "formulas": {"type": "array", "description": "List of {'cell': 'C2', 'formula': '=SUM(A1:A10)'}"},
                    "add_chart": {"type": "object", "description": "{'type': 'bar'|'line'|'pie', 'title': str, 'categories': 'A1:A5', 'data': 'B1:B5', 'anchor': 'F2'}"},
                    "new_sheet": {"type": "string", "description": "Name of a new sheet to create"},
                    "delete_sheet": {"type": "string", "description": "Name of an existing sheet to DELETE (complex sheet removal)."},
                    "rows": {"type": "array", "description": "Complex row operations. Each item: {'action': 'insert'|'delete', 'at': <1-based row number>}. 'insert' pushes a blank row in above row 'at'; 'delete' removes row 'at'."},
                    "columns": {"type": "array", "description": "Complex column operations. Each item: {'action': 'insert'|'delete', 'at': <column letter like 'B'>}. 'insert' puts a blank column before 'at'; 'delete' removes that column."},
                    "clear": {"type": "array", "description": "Ranges to CLEAR (values + formatting), e.g. ['A1:C5', 'F10']. A ':' range clears the rectangle; a bare cell ref clears one cell."},
                    "styles": {"type": "array", "description": "Cell styling. Each item: {'cell': 'A1', 'bold': bool, 'italic': bool, 'size': pt, 'font': 'Calibri', 'fill': 'RRGGBB'}. Use to make headers bold, add fill colors, adjust font sizes."},
                    "inplace": {"type": "boolean", "description": "If true, write back over the original file (no '_edited' copy). Use for the selected target file so no duplicate appears. Default false."},
                },
                "required": ["path"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "edit_presentation",
            "description": "Edit an existing PowerPoint (.pptx). Three modes. (1) APPEND: pass 'add_slides' to stack new slides onto the deck. (2) REBUILD (the right choice when the user asks to REWRITE/change every slide's content/template to match a new topic): pass 'rebuild_slides' to REPLACE every existing slide with this fresh deck. (3) SURGICAL (change only specific existing slides — rename a title, fix text on one slide, add a textbox/table/chart to a slide, or DELETE specific slides — WITHOUT touching the rest): pass 'edit_slides' (per-slide edits) and/or 'remove_slides' (1-based indices of slides to delete). Tell the user which mode (append/rebuild/edit_slides) you used.",
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "Path to existing .pptx file"},
                    "add_slides": {"type": "array", "description": "Slides to APPEND to the existing deck (same shape as create_presentation slides) — use only to add, not to change existing slides."},
                    "rebuild_slides": {"type": "array", "description": "Slides that REPLACE the entire existing deck (same shape as create_presentation slides). Use when the user asks to change every slide / retheme / rewrite the presentation to match a topic."},
                    "edit_slides": {"type": "array", "description": "SURGICAL edits to SPECIFIC existing slides (1-based). Each item: {'slide': 1-based index, 'title': str (new title for that slide), 'replace': {old_text: new_text} (find/replace text on that slide, including inside table cells), 'set_bullets': [str] (replace the body bullets), 'add_textbox': {'text': str, 'top': inches, 'left': inches} (add a text box TO that slide), 'add_table': {'headers': [...], 'rows': [[...]]} (add a table to that slide), 'add_chart': {'type': 'bar'|'line'|'pie', 'title': str, 'categories': [...], 'data': [...]} (add a chart to that slide)}. Use this instead of rebuild when only SOME slides change."},
                    "remove_slides": {"type": "array", "description": "List of 1-based slide indices to DELETE from the deck (surgical 'delete slide N', applied to the ORIGINAL deck numbering). Use instead of rebuild when the user wants one or more specific slides removed."},
                    "inplace": {"type": "boolean", "description": "If true, write back over the original file (no '_edited' copy). Use for the selected target file so no duplicate appears. Default false."},
                },
                "required": ["path"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "read_file",
            "description": "Read the contents of a Word, Excel, PowerPoint, PDF, or plain-text (.txt/.md) file.",
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "Path to the file to read"},
                },
                "required": ["path"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "summarize_file",
            "description": "Summarize a Word, Excel, PowerPoint, PDF, or plain-text (.txt/.md) file (structure + preview).",
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "Path to the file to summarize"},
                },
                "required": ["path"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "get_paragraphs",
            "description": "List the non-empty paragraphs / headings of a Word (.docx) document, most useful to see exactly which section a 'delete page' or 'delete_range' would remove BEFORE doing a destructive delete. Pair this with an anchor-confirmation question to the user: show what the delete will remove and ask before executing.",
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "Path to the existing .docx file"},
                    "limit": {"type": "integer", "description": "Max paragraphs to return (default 100)"},
                },
                "required": ["path"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "edit_text",
            "description": "Edit a plain-text or PDF file. For .txt/.md/.rst: full find/replace/insert/delete with 'edits'. For .pdf: best-effort text find/replace/delete by rewriting the content streams (safe for simple/auto-generated PDFs; layout-heavy or scanned PDFs may not edit cleanly). If the edited path is the user's selected target file, pass inplace=true so no '_edited' duplicate is created.",
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "Path to the .txt/.md/.rst/.log/PDF file to edit"},
                    "edits": {"type": "array", "description": "List of edits. Each edit: {'action': 'replace'|'delete'|'insert_after'|'insert_before'|'delete_range', 'match': exact text to find, 'text': replacement/inserted content, 'end_match': for delete_range, the last text of the block to remove (omit to delete to the end of the file)}. For text files 'replace' swaps 'match' with 'text' (count defaults to 1; use count=-1 for all occurrences)."},
                    "append": {"type": "string", "description": "Text to append at the end of the file"},
                    "inplace": {"type": "boolean", "description": "If true, write back over the original file (no '_edited' copy). Use when editing the selected target file so no duplicate appears. Default false."}
                },
                "required": ["path"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "list_outputs",
            "description": "List files in the outputs folder.",
            "parameters": {
                "type": "object",
                "properties": {},
            },
        },
    },
]


class ToolExecutor:
    def __init__(self, outputs_dir: Path, strict: bool = False):
        self.registry = Registry()
        self.outputs_dir = Path(outputs_dir)
        self.outputs_dir.mkdir(parents=True, exist_ok=True)
        # strict=True surfaces UNEXPECTED executor errors verbatim instead of
        # silently falling through to the file-based editor / masking the cause.
        self.strict = bool(strict)

    def _unexpected(self, exc: Exception, context: str) -> dict:
        """Rethrow-or-return for an unexpected live-editing failure.

        In strict mode the real exception propagates (caught by execute() and
        shown verbatim). In normal mode we return a clear diagnostic so the user
        isn't shown a confusing file-lock error that hides the true cause.
        """
        if self.strict:
            raise exc
        return {
            "error": (
                f"Live edit hit an unexpected error while {context} — I did NOT "
                f"silently fall back and the file was left untouched. "
                f"{type(exc).__name__}: {exc}"
            )
        }

    def _resolve_path(self, path: str) -> str:
        """If the path has no directory part, place it in outputs/."""
        p = Path(path)
        if not p.parent or p.parent == Path("."):
            p = self.outputs_dir / p
        # create parent dirs
        p.parent.mkdir(parents=True, exist_ok=True)
        return str(p)

    def execute(self, name: str, arguments: dict) -> dict:
        handler = getattr(self, f"_{name}", None)
        if handler is None:
            return {"error": f"Unknown tool: {name}"}
        try:
            return handler(arguments)
        except Exception as exc:  # noqa: BLE001
            return {"error": f"{type(exc).__name__}: {exc}"}

    # -- handlers ----------------------------------------------------------
    def _create_document(self, a):
        path = self._resolve_path(a["path"])
        created = self.registry.word.create_document(path, title=a.get("title", ""), body=a.get("body"))
        return {"ok": True, "path": created}

    def _create_workbook(self, a):
        path = self._resolve_path(a["path"])
        created = self.registry.excel.create_workbook(path, sheet=a.get("sheet", "Sheet1"), data=a.get("data"))
        return {"ok": True, "path": created}

    def _create_presentation(self, a):
        path = self._resolve_path(a["path"])
        created = self.registry.ppt.create_presentation(path, slides=a.get("slides"))
        return {"ok": True, "path": created}

    def _edit_document(self, a):
        src = self._resolve_path(a["path"])
        do_inplace = bool(a.get("inplace", False))
        edits = a.get("edits") or []
        append_text = a.get("append")
        content = a.get("content") or []
        del_page = a.get("page")
        confirmed = bool(a.get("confirm", False))
        remove_blank_pages = bool(a.get("remove_blank_pages", False))

        # DESTRUCTIVE-CONFIRMATION BACKSTOP: deleting (page / action 'delete' /
        # 'delete_range') is irreversible. If the model asks for a delete without
        # the user having confirmed a specific anchor, refuse and hand back the
        # document's paragraphs so the model (and user) can pick exactly what is
        # removed. Set confirm=true only AFTER the user confirms the anchor.
        destructive = del_page is not None or any(
            (e or {}).get("action") in ("delete", "delete_range")
            for e in edits
        )
        if destructive and not confirmed:
            try:
                paras = self._get_paragraphs({"path": src, "limit": 60}) or {}
                cand = paras.get("paragraphs", [])
            except Exception:  # noqa: BLE001
                cand = []
            hint = "\n".join(f"  - {p}" for p in cand[:60]) if cand else "  (no readable paragraphs)"
            return {
                "need_confirmation": True,
                "error": (
                    "This is a permanent, irreversible delete. Before removing "
                    "anything I need you to confirm WHICH section/page. See the "
                    f"paragraphs below — tell me the exact one to remove "
                    f"(or provide 'confirm': true after the user agrees).\n"
                    f"Document paragraphs:\n{hint}"
                ),
            }

        # LIVE path: if the .docx is currently open in Word, drive Word via COM
        # so ANY change (replace/insert/delete/append/delete-page) lands in the
        # live open document — no file-lock failure, no duplicate. Works like a
        # coding CLI editing the file you have open.
        try:
            open_in_word = self._is_open_in_word_retry(src)
        except Exception:  # noqa: BLE001
            open_in_word = False
        if open_in_word:
            try:
                return self.registry.word.live_edit(
                    src, edits=edits, append=append_text, page=del_page,
                    remove_empty_paragraphs=remove_blank_pages, content=content,
                    table_edits=a.get("table_edits"))
            except MatchNotFoundError as exc:
                # The text/page the user asked to delete didn't match. This is a
                # user-facing problem, NOT a lock: report clearly and help the
                # model pick a real anchor. NEVER fall through to the locked file.
                return self._live_not_found(src, str(exc))
            except WordNotAvailable:
                pass  # doc isn't actually open in Word -> fall back to file edit
            except Exception as exc:  # noqa: BLE001
                # Unexpected COM / live-edit failure: surface it instead of
                # silently falling through to the locked-file editor (which would
                # mask the real cause with a bogus "file is open" error).
                return self._unexpected(exc, "the live Word edit")

        # A numeric "delete page N" is inherently a live operation: the file-based
        # editor works on the .docx archive and has no notion of rendered pages.
        # So when we can't reach Word live, tell the user instead of silently
        # falling back (which would only hit the lock / do nothing useful).
        if del_page is not None:
            return {
                "error": (
                    "'delete page' needs the document open in the live Word "
                    "window (so Varan can see rendered page boundaries). The file "
                    f"{os.path.basename(src)} is not reachable in running Word "
                    "right now. Open it in Word and try again, or delete a "
                    "section by its exact text with 'delete_range'."
                )
            }

        if remove_blank_pages:
            try:
                result = self.registry.word.remove_empty_paragraphs(src, inplace=do_inplace)
                return {"ok": True, "path": result["path"],
                        "removed": result.get("removed", 0),
                        "message": f"Removed {result.get('removed', 0)} empty paragraph(s) (blank pages)."}
            except (PermissionError, OSError) as exc:
                return {"error": ("The file is currently open and locked by another "
                                  "application (e.g. Word). Close it, or let Varan edit "
                                  "it live through Word instead.")}

        # STRUCTURED CONTENT: appending styled blocks (headings, bullets, italic,
        # Courier, scene breaks) to an existing document. Rendered as REAL Word
        # paragraphs — never plain text / literal '##' markdown.
        if content:
            try:
                result = self.registry.word.insert_content(
                    src, content=content, inplace=do_inplace)
                return {"ok": True, "path": result["path"],
                        "blocks": result.get("blocks", 0),
                        "message": ("Appended {blocks} styled paragraph(s) as real Word "
                                    "content (headings, bullets, italic, fonts)."
                                    .format(blocks=result.get("blocks", 0)))}
            except (PermissionError, OSError) as exc:
                return {"error": ("The file is currently open and locked by another "
                                  "application (e.g. Word). Close it, or let Varan edit "
                                  "it live through Word instead.")}

        # FILE path. delete_range has a native file-backed implementation (it is
        # NOT a live-only operation): split it out, apply each range delete in
        # order, then hand any remaining operations to the standard editor.
        try:
            range_edits = [e for e in edits if (e or {}).get("action") == "delete_range"]
            plain_edits = [e for e in edits if (e or {}).get("action") != "delete_range"]
            cur_path = src
            for ren in range_edits:
                result = self.registry.word.delete_range(
                    cur_path,
                    match=str(ren.get("match", "") or ""),
                    end_match=ren.get("end_match"),
                    end_level=ren.get("end_level"),
                    inplace=do_inplace,
                )
                cur_path = Path(result["path"])
            if range_edits and not plain_edits and append_text is None and not content:
                return {"ok": True, "path": str(cur_path), "deleted": len(range_edits)}
            if range_edits:
                src = cur_path
            edited = self.registry.word.edit_document(
                src,
                edits=plain_edits,
                append=append_text,
                track_changes=bool(a.get("track_changes", False)),
                table_edits=a.get("table_edits"),
                author=a.get("author", "Varan"),
                inplace=do_inplace,
            )
        except (PermissionError, OSError) as exc:
            # The file is open/locked by another program (e.g. Word/Excel).
            return {
                "error": ("The file is currently open and locked by another "
                          "application (e.g. Word). Close it, or let Varan edit "
                          "it live through Word instead."),
            }
        return {"ok": True, "path": edited}

    def _is_open_in_word_retry(self, src):
        """Check if Word has the doc open, retrying once to ride over the
        transient COM init/state hiccup seen when a prior COM handle was used
        and released in the same process (CoInitialize/CoUninitialize churn)."""
        try:
            return self.registry.word.is_open_in_word(src)
        except Exception:  # noqa: BLE001
            pass
        try:
            return self.registry.word.is_open_in_word(src)
        except Exception:  # noqa: BLE001
            return False

    def _live_not_found(self, src, detail: str):
        """Return a clear, actionable message when a live text anchor doesn't
        match, plus the document's real headings so the model can retry with an
        actual anchor instead of guessing."""
        headings = self.registry.word.get_headings(src, limit=30)
        hint = "\n".join(f"  - {h}" for h in headings) if headings else "  (no readable text)"
        return {
            "error": (
                f"{detail}\n\n"
                "This isn't a file-lock problem — the text to edit/delete wasn't "
                "found in the open document. To delete a page or section, use the "
                "exact text from the document below (use action 'delete_range' "
                "with 'match' = first line and 'end_match' = last line of the "
                f"section).\nCurrent headings/text in the document:\n{hint}"
            )
        }

    def _edit_workbook(self, a):
        src = self._resolve_path(a["path"])

        # LIVE path: if the .xlsx is currently open in Excel, drive Excel via COM
        # so writes/formulas/new-sheets land in the live open workbook — no
        # file-lock failure, no duplicate. Works like a coding CLI.
        if self.registry.excel.is_open_in_excel(src):
            try:
                return self.registry.excel.live_edit(
                    src,
                    sheet=a.get("sheet"),
                    writes=a.get("writes"),
                    formulas=a.get("formulas"),
                    new_sheet=a.get("new_sheet"),
                    delete_sheet=a.get("delete_sheet"),
                    rows=a.get("rows"),
                    columns=a.get("columns"),
                    clear=a.get("clear"),
                    styles=a.get("styles"),
                    add_chart=a.get("add_chart"),
                )
            except ExcelNotAvailable:
                pass  # not actually open live -> fall back to file-based edit
            except Exception as exc:  # noqa: BLE001
                return self._unexpected(exc, "the live Excel edit")

        try:
            edited = self.registry.excel.edit_workbook(
                src,
                sheet=a.get("sheet"),
                writes=a.get("writes"),
                formulas=a.get("formulas"),
                add_chart=a.get("add_chart"),
                new_sheet=a.get("new_sheet"),
                delete_sheet=a.get("delete_sheet"),
                rows=a.get("rows"),
                columns=a.get("columns"),
                clear=a.get("clear"),
                styles=a.get("styles"),
                inplace=bool(a.get("inplace", False)),
            )
        except (PermissionError, OSError) as exc:
            return {
                "error": ("The file is currently open and locked by another "
                          "application (e.g. Excel). Close it, or let Varan edit "
                          "it live through Excel instead."),
            }
        return {"ok": True, "path": edited}

    def _edit_presentation(self, a):
        src = self._resolve_path(a["path"])
        add_slides = a.get("add_slides")
        rebuild_slides = a.get("rebuild_slides")
        edit_slides = a.get("edit_slides")
        remove_slides = a.get("remove_slides")
        inplace = bool(a.get("inplace", False))

        # LIVE path: if the .pptx is currently open in PowerPoint, drive
        # PowerPoint via COM so new slides land in the live open presentation —
        # no file-lock failure, no duplicate.
        if self.registry.ppt.is_open_in_ppt(src):
            try:
                return self.registry.ppt.live_edit(
                    src, add_slides=add_slides, rebuild_slides=rebuild_slides,
                    edit_slides=edit_slides, remove_slides=remove_slides)
            except PptNotAvailable:
                pass  # not actually open live -> fall back to file-based edit
            except FillerContentError as exc:
                # Refused degenerate placeholder content on the live path —
                # surface clearly so the model can retry with real content,
                # rather than falling through to the locked file editor.
                return {"error": str(exc)}
            except Exception as exc:  # noqa: BLE001
                return self._unexpected(exc, "the live PowerPoint edit")

        try:
            edited, summary = self.registry.ppt.edit_presentation(
                src, add_slides=add_slides, rebuild_slides=rebuild_slides,
                edit_slides=edit_slides, remove_slides=remove_slides,
                inplace=inplace)
        except (PermissionError, OSError) as exc:
            return {
                "error": ("The file is currently open and locked by another "
                          "application (e.g. PowerPoint). Close it, or let Varan "
                          "edit it live through PowerPoint instead."),
            }
        return {"ok": True, "path": edited, "summary": summary}

    def _edit_text(self, a):
        src = self._resolve_path(a["path"])
        do_inplace = bool(a.get("inplace", False))
        edits = a.get("edits") or []
        append_text = a.get("append")
        kind = Registry.classify(src)
        try:
            if kind == "pdf":
                edited = self.registry.text.edit_pdf(src, edits=edits, inplace=do_inplace)
            else:
                edited = self.registry.text.edit(src, edits=edits, append=append_text,
                                                 inplace=do_inplace)
        except (PermissionError, OSError) as exc:
            return {
                "error": ("The file is currently open and locked by another "
                          "application. Close it and try again."),
            }
        except LookupError as exc:
            headings = self.registry.text.get_headings(src, limit=30)
            hint = "\n".join(f"  - {h}" for h in headings) if headings else "  (no readable text)"
            return {
                "error": (
                    f"{exc}\n\nThis isn't a file-lock problem — the text to "
                    "edit/delete wasn't found. Use the exact text from the file "
                    "below (for a section, use 'delete_range' with 'match' = "
                    f"first line, 'end_match' = last line).\nText in the file:\n{hint}"
                )
            }
        return {"ok": True, "path": edited}

    def _read_file(self, a):
        path = self._resolve_path(a["path"])
        data = self.registry.read(path)
        return {"ok": True, "data": data}

    def _summarize_file(self, a):
        path = self._resolve_path(a["path"])
        data = self.registry.summarize(path)
        return {"ok": True, "summary": data}

    def _get_paragraphs(self, a):
        src = self._resolve_path(a["path"])
        limit = int(a.get("limit", 100) or 100)
        # If the .docx is open live in Word, read its real headings/paragraphs
        # from the live session (best, matches what a delete would remove).
        if self._is_open_in_word_retry(src):
            try:
                return {"ok": True, "paragraphs": self.registry.word.get_headings(src, limit=limit)}
            except Exception:  # noqa: BLE001
                pass
        # Fall back to reading the file on disk:
        #  - .docx via python-docx (paragraph texts, headings preserved)
        #  - text/pdf via the text editor's get_headings
        low = str(src).lower()
        try:
            if low.endswith((".docx", ".doc")):
                return {"ok": True, "paragraphs": self._docx_paragraphs(src, limit)}
            return {"ok": True, "paragraphs": self.registry.text.get_headings(src, limit=limit)}
        except Exception as exc:  # noqa: BLE001
            return {"error": f"{type(exc).__name__}: {exc}"}

    @staticmethod
    def _docx_paragraphs(path: str, limit: int) -> list[str]:
        from docx import Document as D
        d = D(path)
        out: list[str] = []
        empty_count = 0
        empty_styles: dict[str, int] = {}
        collected = 0
        for p in d.paragraphs:
            t = p.text.strip()
            if t:
                if collected < limit:
                    style = (p.style.name or "") if p.style is not None else ""
                    out.append(f"[{style}] {t}" if style else t)
                    collected += 1
            else:
                empty_count += 1
                sname = (p.style.name or "") if p.style is not None else ""
                if sname:
                    empty_styles[sname] = empty_styles.get(sname, 0) + 1
        if empty_count:
            breakdown = ", ".join(f"{n}x {s}" for s, n in sorted(empty_styles.items(), key=lambda kv: -kv[1]))
            summary = f"--- {empty_count} empty paragraph(s)"
            if breakdown:
                summary += f" ({breakdown})"
            summary += " (these are blank pages / page-break spacers — use 'remove_blank_pages' to delete them)"
            out.append(summary)
        return out

    def _list_outputs(self, a):
        files = []
        if self.outputs_dir.exists():
            for f in sorted(self.outputs_dir.iterdir()):
                files.append(str(f.name))
        return {"ok": True, "files": files}
